from pathlib import Path

from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.util import Inches, Pt


# ---------- Theme ----------
W = 13.333
H = 7.5

COLORS = {
    "navy": "122A38",
    "ink": "18313D",
    "teal": "16747E",
    "teal_dark": "0E5962",
    "coral": "E56B4A",
    "mustard": "E6B34A",
    "green": "4F9B78",
    "cream": "F6F1E8",
    "paper": "FFFDFC",
    "pale_teal": "DCEBEA",
    "pale_coral": "F8E3DC",
    "pale_yellow": "F8EDCF",
    "pale_green": "DDEDE4",
    "muted": "667983",
    "line": "CDD8D9",
    "white": "FFFFFF",
    "red": "B84B3B",
}


def rgb(hex_color: str) -> RGBColor:
    h = hex_color.strip("#")
    return RGBColor(int(h[0:2], 16), int(h[2:4], 16), int(h[4:6], 16))


def set_bg(slide, color: str):
    fill = slide.background.fill
    fill.solid()
    fill.fore_color.rgb = rgb(COLORS.get(color, color))


def add_rect(slide, x, y, w, h, fill, radius=True, line=None, line_width=1):
    shape_type = MSO_SHAPE.ROUNDED_RECTANGLE if radius else MSO_SHAPE.RECTANGLE
    shp = slide.shapes.add_shape(shape_type, Inches(x), Inches(y), Inches(w), Inches(h))
    shp.fill.solid()
    shp.fill.fore_color.rgb = rgb(COLORS.get(fill, fill))
    if line:
        shp.line.color.rgb = rgb(COLORS.get(line, line))
        shp.line.width = Pt(line_width)
    else:
        shp.line.fill.background()
    return shp


def add_circle(slide, x, y, d, fill, line=None, line_width=1):
    shp = slide.shapes.add_shape(MSO_SHAPE.OVAL, Inches(x), Inches(y), Inches(d), Inches(d))
    shp.fill.solid()
    shp.fill.fore_color.rgb = rgb(COLORS.get(fill, fill))
    if line:
        shp.line.color.rgb = rgb(COLORS.get(line, line))
        shp.line.width = Pt(line_width)
    else:
        shp.line.fill.background()
    return shp


def add_text(slide, text, x, y, w, h, size=20, color="ink", bold=False,
             font="Aptos", align=PP_ALIGN.LEFT, valign=MSO_ANCHOR.TOP,
             margin=0.05, line_spacing=1.0, italic=False):
    box = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    tf = box.text_frame
    tf.clear()
    tf.word_wrap = True
    tf.margin_left = Inches(margin)
    tf.margin_right = Inches(margin)
    tf.margin_top = Inches(margin)
    tf.margin_bottom = Inches(margin)
    tf.vertical_anchor = valign
    p = tf.paragraphs[0]
    p.text = text
    p.alignment = align
    p.line_spacing = line_spacing
    p.font.name = font
    p.font.size = Pt(size)
    p.font.bold = bold
    p.font.italic = italic
    p.font.color.rgb = rgb(COLORS.get(color, color))
    return box


def add_rich_text(slide, runs, x, y, w, h, size=20, align=PP_ALIGN.LEFT,
                  valign=MSO_ANCHOR.TOP, margin=0.05, line_spacing=1.0):
    box = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    tf = box.text_frame
    tf.clear()
    tf.word_wrap = True
    tf.margin_left = Inches(margin)
    tf.margin_right = Inches(margin)
    tf.margin_top = Inches(margin)
    tf.margin_bottom = Inches(margin)
    tf.vertical_anchor = valign
    p = tf.paragraphs[0]
    p.alignment = align
    p.line_spacing = line_spacing
    for item in runs:
        r = p.add_run()
        r.text = item[0]
        r.font.name = item[4] if len(item) > 4 else "Aptos"
        r.font.size = Pt(item[3] if len(item) > 3 else size)
        r.font.bold = item[1]
        r.font.color.rgb = rgb(COLORS.get(item[2], item[2]))
    return box


def add_pill(slide, text, x, y, w, fill="pale_teal", color="teal_dark", size=12):
    add_rect(slide, x, y, w, 0.38, fill, radius=True)
    add_text(slide, text, x + 0.05, y + 0.02, w - 0.1, 0.30, size=size,
             color=color, bold=True, align=PP_ALIGN.CENTER, valign=MSO_ANCHOR.MIDDLE, margin=0)


def add_header(slide, number, title, kicker=None):
    add_text(slide, f"{number:02d}", 0.55, 0.33, 0.55, 0.38, size=13, color="coral", bold=True,
             align=PP_ALIGN.CENTER, valign=MSO_ANCHOR.MIDDLE, margin=0)
    add_rect(slide, 1.18, 0.45, 0.32, 0.05, "coral", radius=False)
    if kicker:
        add_text(slide, kicker.upper(), 1.62, 0.28, 5.0, 0.27, size=10, color="teal", bold=True, margin=0)
        add_text(slide, title, 1.60, 0.55, 10.8, 0.60, size=27, color="ink", bold=True, margin=0)
    else:
        add_text(slide, title, 1.60, 0.32, 10.8, 0.62, size=27, color="ink", bold=True, margin=0)
    add_rect(slide, 0.55, 1.18, 12.2, 0.025, "line", radius=False)


def add_footer(slide, number, dark=False):
    c = "white" if dark else "muted"
    add_text(slide, "РУССКИЙ ЯЗЫК • ЛИТЕРАТУРА • ОЛИМПИАДА", 0.58, 7.08, 4.6, 0.22,
             size=8, color=c, bold=True, margin=0)
    add_text(slide, str(number), 12.25, 7.04, 0.45, 0.25, size=9, color=c, bold=True,
             align=PP_ALIGN.RIGHT, margin=0)


def add_number_icon(slide, n, x, y, fill="teal", color="white", d=0.52):
    add_circle(slide, x, y, d, fill)
    add_text(slide, str(n), x, y + 0.01, d, d - 0.02, size=14, color=color, bold=True,
             align=PP_ALIGN.CENTER, valign=MSO_ANCHOR.MIDDLE, margin=0)


def add_card(slide, x, y, w, h, title, body, accent="teal", fill="paper", icon=None):
    add_rect(slide, x, y, w, h, fill, radius=True, line="line", line_width=0.8)
    add_rect(slide, x, y, 0.08, h, accent, radius=False)
    if icon:
        add_circle(slide, x + 0.27, y + 0.27, 0.48, accent)
        add_text(slide, icon, x + 0.27, y + 0.27, 0.48, 0.48, size=16, color="white", bold=True,
                 align=PP_ALIGN.CENTER, valign=MSO_ANCHOR.MIDDLE, margin=0)
        tx = x + 0.90
        tw = w - 1.12
    else:
        tx = x + 0.30
        tw = w - 0.55
    add_text(slide, title, tx, y + 0.22, tw, 0.42, size=16, color="ink", bold=True, margin=0)
    add_text(slide, body, x + 0.30, y + 0.78, w - 0.58, h - 0.96, size=11.5, color="muted",
             margin=0, line_spacing=1.05)


def add_bullet_list(slide, items, x, y, w, h, color="ink", accent="coral", size=16, gap=0.52):
    for i, item in enumerate(items):
        yy = y + i * gap
        add_circle(slide, x, yy + 0.10, 0.13, accent)
        add_text(slide, item, x + 0.27, yy, w - 0.27, gap, size=size, color=color,
                 valign=MSO_ANCHOR.TOP, margin=0)


def build_presentation():
    prs = Presentation()
    prs.slide_width = Inches(W)
    prs.slide_height = Inches(H)
    blank = prs.slide_layouts[6]

    # 1. Cover
    slide = prs.slides.add_slide(blank)
    set_bg(slide, "navy")
    add_pill(slide, "МАСТЕР-КЛАСС ДЛЯ ПЕДАГОГОВ", 0.70, 0.55, 2.75,
             fill="coral", color="white", size=10)
    add_text(slide, "Эффективные методы\nподготовки к олимпиадам", 0.72, 1.32, 7.35, 1.75,
             size=31, color="white", bold=True, margin=0, line_spacing=0.92)
    add_text(slide, "по русскому языку и литературе", 0.73, 3.24, 7.0, 0.58,
             size=22, color="mustard", bold=True, margin=0)
    add_text(slide, "Опыт  •  анализ  •  современные технологии", 0.74, 4.12, 6.6, 0.36,
             size=13, color="pale_teal", bold=True, margin=0)
    add_rect(slide, 0.74, 4.83, 6.05, 0.03, "teal", radius=False)
    add_text(slide, "40 минут  |  10 слайдов  |  практические задания", 0.75, 5.07, 6.3, 0.34,
             size=11, color="white", margin=0)
    add_text(slide, "Автор: ____________________\nОрганизация: ____________________", 0.75, 5.70, 5.5, 0.75,
             size=11.5, color="pale_teal", margin=0, line_spacing=1.2)

    # visual open book
    add_circle(slide, 9.10, 0.92, 3.48, "teal_dark")
    left_page = add_rect(slide, 8.40, 2.03, 2.02, 2.75, "paper", radius=True)
    left_page.rotation = -7
    right_page = add_rect(slide, 10.33, 2.03, 2.02, 2.75, "paper", radius=True)
    right_page.rotation = 7
    add_rect(slide, 10.26, 2.24, 0.06, 2.35, "mustard", radius=False)
    for i, width in enumerate([1.22, 1.42, 1.10, 1.34]):
        add_rect(slide, 8.78, 2.50 + i * 0.42, width, 0.04, "line", radius=False)
    for i, width in enumerate([1.28, 1.06, 1.40, 1.20]):
        add_rect(slide, 10.72, 2.50 + i * 0.42, width, 0.04, "line", radius=False)
    add_pill(slide, "СЛОВО", 8.25, 5.26, 1.10, fill="pale_coral", color="coral", size=10)
    add_pill(slide, "СМЫСЛ", 9.65, 5.60, 1.18, fill="pale_yellow", color="ink", size=10)
    add_pill(slide, "ДОКАЗАТЬ", 11.05, 5.18, 1.48, fill="pale_green", color="green", size=10)
    add_footer(slide, 1, dark=True)

    # 2. What the olympiad tests
    slide = prs.slides.add_slide(blank)
    set_bg(slide, "cream")
    add_header(slide, 2, "Что на самом деле проверяет олимпиада?", "Не набор правил")
    add_text(slide, "ЗНАТЬ", 0.78, 1.55, 2.05, 0.45, size=16, color="muted", bold=True,
             align=PP_ALIGN.CENTER, margin=0)
    add_text(slide, "≠", 2.73, 1.47, 0.72, 0.52, size=28, color="coral", bold=True,
             align=PP_ALIGN.CENTER, margin=0)
    add_text(slide, "ПРИМЕНЯТЬ", 3.48, 1.55, 2.45, 0.45, size=16, color="teal", bold=True,
             align=PP_ALIGN.CENTER, margin=0)
    add_text(slide, "Олимпиадный результат появляется, когда ученик выполняет четыре действия:",
             6.35, 1.50, 5.92, 0.55, size=14, color="ink", bold=True, margin=0)

    cards = [
        ("1", "Увидеть закономерность", "Сравнить примеры и заметить, что именно меняется.", "teal", "pale_teal"),
        ("2", "Выдвинуть гипотезу", "Сформулировать возможное объяснение своими словами.", "coral", "pale_coral"),
        ("3", "Доказать", "Подтвердить ответ признаками, цитатой или контрпримером.", "mustard", "pale_yellow"),
        ("4", "Точно оформить", "Дать полный ответ без лишних рассуждений и двусмысленности.", "green", "pale_green"),
    ]
    positions = [(0.78, 2.35), (3.92, 2.35), (7.06, 2.35), (10.20, 2.35)]
    for (num, title, body, accent, fill), (x, y) in zip(cards, positions):
        add_rect(slide, x, y, 2.38, 2.78, fill, radius=True)
        add_circle(slide, x + 0.22, y + 0.22, 0.58, accent)
        add_text(slide, num, x + 0.22, y + 0.22, 0.58, 0.58, size=17, color="white", bold=True,
                 align=PP_ALIGN.CENTER, valign=MSO_ANCHOR.MIDDLE, margin=0)
        add_text(slide, title, x + 0.22, y + 1.03, 1.94, 0.76, size=17, color="ink", bold=True, margin=0)
        add_text(slide, body, x + 0.22, y + 1.87, 1.94, 0.72, size=11, color="muted", margin=0)

    add_rect(slide, 0.78, 5.55, 11.80, 0.82, "navy", radius=True)
    add_text(slide, "Фокус подготовки:", 1.05, 5.78, 2.0, 0.28, size=13, color="mustard", bold=True, margin=0)
    add_text(slide, "не «сколько заданий решено», а «как ученик объясняет решение»", 2.83, 5.74, 8.92, 0.34,
             size=17, color="white", bold=True, margin=0)
    add_footer(slide, 2)

    # 3. System
    slide = prs.slides.add_slide(blank)
    set_bg(slide, "paper")
    add_header(slide, 3, "Система подготовки: шесть последовательных шагов", "От отбора до коррекции")
    labels = [
        ("Диагностика", "Что уже умеет?"),
        ("Цель", "Одна зона роста"),
        ("Тренировка", "Приём + практика"),
        ("Объяснение", "Почему так?"),
        ("Пробная работа", "Время и критерии"),
        ("Коррекция", "Следующее действие"),
    ]
    xs = [0.72, 2.78, 4.84, 6.90, 8.96, 11.02]
    accents = ["teal", "coral", "mustard", "green", "teal_dark", "coral"]
    for i, ((title, sub), x, accent) in enumerate(zip(labels, xs, accents), start=1):
        if i < 6:
            add_rect(slide, x + 0.77, 2.53, 1.28, 0.05, "line", radius=False)
        add_circle(slide, x, 2.03, 0.92, accent)
        add_text(slide, str(i), x, 2.04, 0.92, 0.90, size=22, color="white", bold=True,
                 align=PP_ALIGN.CENTER, valign=MSO_ANCHOR.MIDDLE, margin=0)
        add_text(slide, title, x - 0.34, 3.16, 1.62, 0.50, size=14, color="ink", bold=True,
                 align=PP_ALIGN.CENTER, margin=0)
        add_text(slide, sub, x - 0.38, 3.73, 1.70, 0.46, size=10.5, color="muted",
                 align=PP_ALIGN.CENTER, margin=0)

    add_rect(slide, 0.78, 4.78, 5.56, 1.30, "pale_teal", radius=True)
    add_text(slide, "Индивидуальный маршрут", 1.08, 5.03, 2.70, 0.38, size=17, color="teal_dark", bold=True, margin=0)
    add_text(slide, "сильная сторона  +  одна измеримая зона роста", 1.08, 5.49, 4.85, 0.34,
             size=13, color="ink", margin=0)
    add_rect(slide, 6.63, 4.78, 5.92, 1.30, "pale_coral", radius=True)
    add_text(slide, "Главный принцип", 6.94, 5.03, 2.15, 0.38, size=17, color="coral", bold=True, margin=0)
    add_text(slide, "каждая ошибка заканчивается конкретным следующим заданием", 6.94, 5.49, 5.20, 0.34,
             size=13, color="ink", margin=0)
    add_footer(slide, 3)

    # 4. Diagnostics
    slide = prs.slides.add_slide(blank)
    set_bg(slide, "cream")
    add_header(slide, 4, "Диагностика за 20 минут", "Не оценка, а карта затруднений")
    add_rect(slide, 0.72, 1.48, 5.86, 3.48, "paper", radius=True, line="line")
    add_pill(slide, "РУССКИЙ ЯЗЫК", 1.03, 1.77, 1.68, fill="pale_teal", color="teal_dark", size=10)
    lang_items = [
        "лексика и значение слова",
        "морфология и формы слов",
        "синтаксис и пунктуация",
        "редактирование и доказательство",
    ]
    add_bullet_list(slide, lang_items, 1.03, 2.38, 4.95, 2.15, size=14.5, accent="teal", gap=0.56)
    add_rect(slide, 6.78, 1.48, 5.84, 3.48, "paper", radius=True, line="line")
    add_pill(slide, "ЛИТЕРАТУРА", 7.09, 1.77, 1.42, fill="pale_coral", color="coral", size=10)
    lit_items = [
        "понимание фактов текста",
        "художественная деталь",
        "средство → эффект → смысл",
        "аргументация цитатой",
    ]
    add_bullet_list(slide, lit_items, 7.09, 2.38, 4.95, 2.15, size=14.5, accent="coral", gap=0.56)

    add_text(slide, "КАРТА ПРОДВИЖЕНИЯ", 0.78, 5.24, 2.43, 0.30, size=11, color="teal", bold=True, margin=0)
    cols = [
        ("Ошибка", "Приём назван, роль не объяснена", 0.78, 3.35, "pale_yellow"),
        ("Причина", "Термин выучен без связи со смыслом", 4.28, 3.75, "pale_coral"),
        ("Следующее действие", "3 фрагмента по формуле «приём → эффект»", 8.18, 4.37, "pale_green"),
    ]
    for title, body, x, w, fill in cols:
        add_rect(slide, x, 5.62, w, 0.96, fill, radius=True)
        add_text(slide, title, x + 0.20, 5.76, w - 0.40, 0.24, size=10.5, color="muted", bold=True, margin=0)
        add_text(slide, body, x + 0.20, 6.08, w - 0.40, 0.34, size=11.5, color="ink", bold=True, margin=0)
    add_footer(slide, 4)

    # 5. Methods
    slide = prs.slides.add_slide(blank)
    set_bg(slide, "paper")
    add_header(slide, 5, "Четыре метода, которые работают", "Коротко, регулярно, с объяснением")
    add_card(slide, 0.72, 1.52, 5.88, 2.14, "Лингвистический детектив",
             "Найти закономерность → сформулировать гипотезу → подтвердить двумя примерами.",
             accent="teal", fill="pale_teal", icon="Л")
    add_card(slide, 6.75, 1.52, 5.86, 2.14, "Редакторская мастерская",
             "Исправить ошибку, назвать её тип и объяснить, почему новый вариант точнее.",
             accent="coral", fill="pale_coral", icon="Р")
    add_card(slide, 0.72, 3.88, 5.88, 2.14, "Читательская лаборатория",
             "Деталь → средство выразительности → эффект → смысл текста.",
             accent="mustard", fill="pale_yellow", icon="Ч")
    add_card(slide, 6.75, 3.88, 5.86, 2.14, "Объясни другому",
             "Ученик проговаривает ход решения. Партнёр задаёт один уточняющий вопрос.",
             accent="green", fill="pale_green", icon="О")
    add_rect(slide, 3.86, 6.30, 5.62, 0.50, "navy", radius=True)
    add_text(slide, "Ритм: 5 минут на каждом уроке + 1 глубокое занятие в неделю", 4.05, 6.41, 5.25, 0.25,
             size=10.5, color="white", bold=True, align=PP_ALIGN.CENTER, margin=0)
    add_footer(slide, 5)

    # 6. Russian language practice
    slide = prs.slides.add_slide(blank)
    set_bg(slide, "cream")
    add_header(slide, 6, "Практика 1. Лингвистический детектив", "Русский язык")
    add_rect(slide, 0.72, 1.50, 5.65, 4.94, "navy", radius=True)
    add_pill(slide, "ЗАДАНИЕ", 1.02, 1.82, 1.15, fill="coral", color="white", size=10)
    add_text(slide, "Стекло заблестело на солнце.\n\nМолоко стекло на пол.", 1.03, 2.42, 4.95, 1.55,
             size=22, color="white", bold=True, margin=0, line_spacing=1.05)
    add_text(slide, "Определите часть речи слова «стекло»\nв каждом предложении. Докажите ответ.",
             1.03, 4.38, 4.83, 0.92, size=15, color="pale_teal", margin=0)
    add_rect(slide, 1.02, 5.67, 4.95, 0.42, "teal_dark", radius=True)
    add_text(slide, "Время на обсуждение: 60 секунд", 1.12, 5.76, 4.75, 0.22, size=10.5,
             color="white", bold=True, align=PP_ALIGN.CENTER, margin=0)

    add_text(slide, "ДОКАЗАТЕЛЬСТВО", 6.83, 1.60, 2.28, 0.30, size=11, color="teal", bold=True, margin=0)
    add_card(slide, 6.82, 2.03, 5.72, 1.72, "1. Существительное",
             "«Стекло» отвечает на вопрос что?, называет предмет, является подлежащим.",
             accent="teal", fill="paper", icon="1")
    add_card(slide, 6.82, 3.96, 5.72, 1.72, "2. Глагол",
             "«Стекло» отвечает на вопрос что сделало?, начальная форма — стечь, является сказуемым.",
             accent="coral", fill="paper", icon="2")
    add_rect(slide, 6.82, 5.93, 5.72, 0.50, "pale_yellow", radius=True)
    add_text(slide, "Вывод: одинаковые формы разных слов — омоформы.", 7.04, 6.05, 5.30, 0.24,
             size=12.5, color="ink", bold=True, margin=0)
    add_footer(slide, 6)

    # 7. Literature practice
    slide = prs.slides.add_slide(blank)
    set_bg(slide, "paper")
    add_header(slide, 7, "Практика 2. Литературная лаборатория", "Литература")
    add_rect(slide, 0.72, 1.49, 6.17, 3.08, "pale_teal", radius=True)
    add_text(slide, "«", 1.05, 1.63, 0.58, 0.60, size=39, color="teal", bold=True, margin=0)
    fragment = ("Ветер настойчиво стучал в окно. На столе лежало неотправленное письмо. "
                "Мальчик прислушивался к шагам на лестнице.")
    add_text(slide, fragment, 1.46, 1.99, 4.85, 1.55, size=18, color="ink", bold=True,
             margin=0, line_spacing=1.12)
    add_text(slide, "Определите настроение, найдите приём и объясните роль письма.",
             1.05, 3.80, 5.30, 0.40, size=12.5, color="teal_dark", bold=True, margin=0)

    add_text(slide, "ФОРМУЛА ОТВЕТА", 7.24, 1.55, 2.20, 0.30, size=11, color="coral", bold=True, margin=0)
    chain = [
        ("ДЕТАЛЬ", "неотправленное письмо", "teal", "pale_teal"),
        ("ПРИЁМ", "«ветер стучал» — олицетворение", "coral", "pale_coral"),
        ("ЭФФЕКТ", "напряжение и беспокойство", "mustard", "pale_yellow"),
        ("СМЫСЛ", "ожидание, нерешительность, недосказанность", "green", "pale_green"),
    ]
    y = 2.02
    for i, (title, body, accent, fill) in enumerate(chain):
        add_rect(slide, 7.22, y, 5.28, 0.78, fill, radius=True)
        add_pill(slide, title, 7.43, y + 0.19, 1.08, fill=accent, color="white", size=9)
        add_text(slide, body, 8.74, y + 0.19, 3.50, 0.38, size=11.5, color="ink", bold=True, margin=0)
        if i < 3:
            add_text(slide, "↓", 9.55, y + 0.77, 0.50, 0.32, size=18, color="muted", bold=True,
                     align=PP_ALIGN.CENTER, margin=0)
        y += 1.04

    add_rect(slide, 0.72, 4.86, 6.17, 1.40, "navy", radius=True)
    add_text(slide, "Критерий полного ответа", 1.02, 5.09, 2.55, 0.32, size=14, color="mustard", bold=True, margin=0)
    add_text(slide, "термин + короткая цитата + объяснение роли + вывод о смысле", 1.02, 5.55, 5.45, 0.38,
             size=14, color="white", bold=True, margin=0)
    add_footer(slide, 7)

    # 8. Technology
    slide = prs.slides.add_slide(blank)
    set_bg(slide, "cream")
    add_header(slide, 8, "Современные технологии: где помогают, где мешают", "Инструмент, а не замена мышлению")
    add_rect(slide, 0.72, 1.52, 5.88, 4.70, "pale_green", radius=True)
    add_pill(slide, "ПОМОГАЮТ", 1.02, 1.84, 1.42, fill="green", color="white", size=10)
    help_items = [
        "онлайн-формы — быстрая диагностика",
        "электронные словари и корпус — проверка употребления",
        "общий документ — совместная разметка текста",
        "ИИ — черновики заданий и варианты формулировок",
    ]
    add_bullet_list(slide, help_items, 1.05, 2.48, 5.05, 2.74, size=13.5, accent="green", gap=0.64)

    add_rect(slide, 6.75, 1.52, 5.86, 4.70, "pale_coral", radius=True)
    add_pill(slide, "МЕШАЮТ", 7.05, 1.84, 1.28, fill="red", color="white", size=10)
    risk_items = [
        "непроверенные ответы и вымышленные примеры",
        "готовая интерпретация вместо чтения текста",
        "тесты без объяснения выбора",
        "загрузка персональных данных учеников",
    ]
    add_bullet_list(slide, risk_items, 7.08, 2.48, 5.05, 2.74, size=13.5, accent="red", gap=0.64)

    add_rect(slide, 2.58, 6.42, 8.18, 0.42, "navy", radius=True)
    add_text(slide, "ИИ создаёт черновик  →  учитель проверяет  →  ученик самостоятельно решает",
             2.75, 6.51, 7.84, 0.22, size=11, color="white", bold=True,
             align=PP_ALIGN.CENTER, margin=0)
    add_footer(slide, 8)

    # 9. Four-week plan
    slide = prs.slides.add_slide(blank)
    set_bg(slide, "paper")
    add_header(slide, 9, "Компактный план подготовки на четыре недели", "Без перегрузки")
    week_data = [
        ("1", "Диагностика", "20 минут", "Выбрать одну сильную сторону и одну зону роста.", "teal", "pale_teal"),
        ("2", "Приёмы", "2 × 25 минут", "Тренировать язык и анализ текста по понятным формулам.", "coral", "pale_coral"),
        ("3", "Пробная работа", "45 минут", "Решить вариант по времени и проверить по критериям.", "mustard", "pale_yellow"),
        ("4", "Коррекция", "2 × 25 минут", "Разобрать причины ошибок и повторить только слабые места.", "green", "pale_green"),
    ]
    xs = [0.72, 3.82, 6.92, 10.02]
    for (num, title, time, body, accent, fill), x in zip(week_data, xs):
        add_rect(slide, x, 1.62, 2.58, 4.46, fill, radius=True)
        add_circle(slide, x + 0.87, 1.95, 0.84, accent)
        add_text(slide, num, x + 0.87, 1.95, 0.84, 0.84, size=20, color="white", bold=True,
                 align=PP_ALIGN.CENTER, valign=MSO_ANCHOR.MIDDLE, margin=0)
        add_text(slide, f"НЕДЕЛЯ {num}", x + 0.35, 2.98, 1.88, 0.24, size=9.5, color=accent, bold=True,
                 align=PP_ALIGN.CENTER, margin=0)
        add_text(slide, title, x + 0.28, 3.31, 2.02, 0.68, size=17, color="ink", bold=True,
                 align=PP_ALIGN.CENTER, margin=0)
        add_pill(slide, time, x + 0.54, 4.14, 1.50, fill="paper", color="ink", size=9.5)
        add_text(slide, body, x + 0.32, 4.82, 1.94, 0.92, size=11, color="muted",
                 align=PP_ALIGN.CENTER, margin=0)

    add_rect(slide, 2.06, 6.35, 9.22, 0.48, "pale_teal", radius=True)
    add_text(slide, "Оптимальный ритм: короткая практика между занятиями + подробный разбор одной работы",
             2.25, 6.46, 8.84, 0.24, size=11.5, color="teal_dark", bold=True,
             align=PP_ALIGN.CENTER, margin=0)
    add_footer(slide, 9)

    # 10. Final formula
    slide = prs.slides.add_slide(blank)
    set_bg(slide, "navy")
    add_pill(slide, "ИТОГ", 0.72, 0.54, 0.92, fill="coral", color="white", size=10)
    add_text(slide, "Формула устойчивого результата", 0.72, 1.18, 8.20, 0.72,
             size=30, color="white", bold=True, margin=0)
    add_text(slide, "Победа начинается не со сложного задания, а с понятной системы работы.",
             0.73, 2.01, 8.10, 0.48, size=15, color="pale_teal", margin=0)

    parts = [
        ("СИСТЕМА", "teal"),
        ("РАССУЖДЕНИЕ", "coral"),
        ("ТЕКСТ", "mustard"),
        ("ОБРАТНАЯ СВЯЗЬ", "green"),
    ]
    x = 0.72
    widths = [1.78, 2.45, 1.38, 2.72]
    for i, ((label, accent), width) in enumerate(zip(parts, widths)):
        add_rect(slide, x, 2.86, width, 0.72, accent, radius=True)
        add_text(slide, label, x + 0.08, 3.07, width - 0.16, 0.28, size=12, color="white", bold=True,
                 align=PP_ALIGN.CENTER, margin=0)
        x += width + 0.22
        if i < 3:
            add_text(slide, "+", x - 0.18, 3.03, 0.20, 0.30, size=17, color="white", bold=True,
                     align=PP_ALIGN.CENTER, margin=0)

    add_rect(slide, 8.95, 0.83, 3.65, 5.85, "paper", radius=True)
    add_text(slide, "ЧЕК-ЛИСТ ПЕДАГОГА", 9.30, 1.25, 2.95, 0.38, size=14, color="teal_dark", bold=True,
             align=PP_ALIGN.CENTER, margin=0)
    checklist = [
        "Есть конкретная цель",
        "Задание требует доказательства",
        "Критерии известны заранее",
        "Ошибка разобрана по причине",
        "Назначено следующее действие",
    ]
    for i, item in enumerate(checklist):
        yy = 1.95 + i * 0.72
        add_circle(slide, 9.35, yy, 0.34, "pale_green", line="green")
        add_text(slide, "✓", 9.35, yy, 0.34, 0.34, size=13, color="green", bold=True,
                 align=PP_ALIGN.CENTER, valign=MSO_ANCHOR.MIDDLE, margin=0)
        add_text(slide, item, 9.86, yy + 0.02, 2.25, 0.34, size=11.7, color="ink", bold=True, margin=0)

    add_rect(slide, 0.72, 4.23, 7.80, 1.10, "teal_dark", radius=True)
    add_text(slide, "Рефлексия: какой один приём вы примените уже на ближайшем занятии?",
             1.02, 4.51, 7.20, 0.44, size=17, color="white", bold=True,
             align=PP_ALIGN.CENTER, margin=0)
    add_text(slide, "Спасибо за работу!", 0.74, 5.90, 4.90, 0.48, size=20, color="mustard", bold=True, margin=0)
    add_text(slide, "Диагностика → цель → тренировка → объяснение → проба → коррекция",
             0.74, 6.45, 7.70, 0.34, size=11.5, color="pale_teal", bold=True, margin=0)
    add_footer(slide, 10, dark=True)

    out_dir = Path(__file__).resolve().parent
    out_path = out_dir / "masterclass_russian_olympiad_10_slides.pptx"
    prs.save(out_path)
    print(f"Saved: {out_path}")
    print(f"Slides: {len(prs.slides)}")


if __name__ == "__main__":
    build_presentation()
